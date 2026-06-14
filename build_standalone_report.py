#!/usr/bin/env python
"""Render a seed scouting report Markdown file as a standalone DOCX.

The document borrows styles, page setup, headers, and footers from
Tournament_Preparation_Manual.docx, but the manual is opened read-only and is
never saved.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "Tournament_Preparation_Manual.docx"
HEADER_OLD = "Tournament Preparation Manual"
BULLET_NUMID = 2
USABLE_WIDTH_DXA = 9360
TABLE_FONT_PT = 9

LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]*\)")


def strip_links(text: str) -> str:
    return LINK_RE.sub(r"\1", text)


def _fmt(bold: bool, italic: bool, mono: bool) -> dict[str, bool]:
    out: dict[str, bool] = {}
    if bold:
        out["bold"] = True
    if italic:
        out["italic"] = True
    if mono:
        out["mono"] = True
    return out


def _emit_inline(text: str, bold: bool, italic: bool, out: list[tuple[str, dict[str, bool]]]) -> None:
    while text:
        match = re.search(r"(\*\*|\*|`)", text)
        if not match:
            out.append((text, _fmt(bold, italic, False)))
            return
        if match.start() > 0:
            out.append((text[: match.start()], _fmt(bold, italic, False)))
        marker = match.group(1)
        rest = text[match.end() :]
        if marker == "`":
            end = rest.find("`")
            if end < 0:
                out.append(("`" + rest, _fmt(bold, italic, False)))
                return
            out.append((rest[:end], _fmt(bold, italic, True)))
            text = rest[end + 1 :]
        elif marker == "**":
            end = rest.find("**")
            if end < 0:
                out.append(("**" + rest, _fmt(bold, italic, False)))
                return
            _emit_inline(rest[:end], True, italic, out)
            text = rest[end + 2 :]
        else:
            end = rest.find("*")
            if end < 0:
                out.append(("*" + rest, _fmt(bold, italic, False)))
                return
            _emit_inline(rest[:end], bold, True, out)
            text = rest[end + 1 :]


def inline_runs(text: str) -> list[tuple[str, dict[str, bool]]]:
    text = strip_links(text).replace("⚠️", "").replace("⚠", "").strip()
    out: list[tuple[str, dict[str, bool]]] = []
    _emit_inline(text, False, False, out)
    return out or [("", {})]


def table_cells(row: str) -> list[str]:
    row = row.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [cell.strip() for cell in row.split("|")]


def is_block_start(line: str) -> bool:
    text = line.strip()
    if not text:
        return True
    if re.match(r"#{1,6}\s", text):
        return True
    if text.startswith(("|", ">")):
        return True
    if re.fullmatch(r"-{3,}", text):
        return True
    if re.match(r"-\s+", text) or re.match(r"\d+\.\s+", text):
        return True
    return False


def collect_continuation(lines: list[str], index: int) -> tuple[str, int]:
    parts: list[str] = []
    while index < len(lines) and not is_block_start(lines[index]):
        parts.append(lines[index].strip())
        index += 1
    return " ".join(part for part in parts if part), index


def is_separator(line: str) -> bool:
    text = line.strip()
    return bool(text) and set(text) <= set("|:- ") and "-" in text


def parse_blocks(markdown: str) -> list[tuple]:
    lines = markdown.split("\n")
    blocks: list[tuple] = []
    index = 0
    while index < len(lines):
        raw = lines[index].rstrip()
        text = raw.strip()
        if not text or re.fullmatch(r"-{3,}", text):
            index += 1
            continue

        heading = re.match(r"(#{1,6})\s+(.*)", raw)
        if heading:
            blocks.append(("h", len(heading.group(1)), heading.group(2).strip()))
            index += 1
            continue

        if raw.lstrip().startswith("|") and index + 1 < len(lines) and is_separator(lines[index + 1]):
            table_lines = [raw]
            separator = lines[index + 1]
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            aligns = [
                "center" if cell.startswith(":") and cell.endswith(":") else "right" if cell.endswith(":") else "left"
                for cell in table_cells(separator)
            ]
            blocks.append(("table", table_cells(table_lines[0]), aligns, [table_cells(row) for row in table_lines[1:]]))
            continue

        if raw.lstrip().startswith(">"):
            quote: list[str] = []
            while index < len(lines) and lines[index].lstrip().startswith(">"):
                quote.append(re.sub(r"^\s*>\s?", "", lines[index]).strip())
                index += 1
            blocks.append(("quote", " ".join(part for part in quote if part)))
            continue

        bullet = re.match(r"\s*-\s+(.*)", raw)
        if bullet:
            extra, index = collect_continuation(lines, index + 1)
            blocks.append(("bullet", (bullet.group(1).strip() + " " + extra).strip()))
            continue

        number = re.match(r"\s*(\d+)\.\s+(.*)", raw)
        if number:
            extra, index = collect_continuation(lines, index + 1)
            blocks.append(("number", number.group(1), (number.group(2).strip() + " " + extra).strip()))
            continue

        extra, index = collect_continuation(lines, index + 1)
        blocks.append(("p", (text + " " + extra).strip()))
    return blocks


def add_runs(paragraph, runs: list[tuple[str, dict[str, bool]]]):
    for text, fmt in runs:
        run = paragraph.add_run(text)
        if fmt.get("bold"):
            run.bold = True
        if fmt.get("italic"):
            run.italic = True
        if fmt.get("mono"):
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")


def set_paragraph_style(paragraph, style_id: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:pStyle"))
    if existing is not None:
        p_pr.remove(existing)
    element = OxmlElement("w:pStyle")
    element.set(qn("w:val"), style_id)
    p_pr.insert(0, element)


def set_bullet(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(BULLET_NUMID))
    num_pr.append(num_id)
    p_pr.append(num_pr)


def set_cell_margins(cell, margin_dxa: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(margin_dxa))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_row_repeats(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def column_widths(headers: list[str]) -> list[int]:
    cols = len(headers)
    normalized = [strip_links(h).strip().lower() for h in headers]
    if cols == 2:
        ratios = [0.5, 0.5]
    elif cols == 3:
        ratios = [0.3, 0.12, 0.58]
    elif cols == 5 and normalized[-1].startswith("suggested"):
        ratios = [0.28, 0.12, 0.10, 0.10, 0.40]
    elif cols == 6:
        ratios = [0.18, 0.10, 0.11, 0.22, 0.13, 0.26]
    elif cols == 7:
        ratios = [0.18, 0.10, 0.10, 0.09, 0.09, 0.10, 0.34]
    else:
        ratios = [1 / cols] * cols
    widths = [int(USABLE_WIDTH_DXA * r) for r in ratios]
    widths[-1] += USABLE_WIDTH_DXA - sum(widths)
    return widths


def style_table(table, headers: list[str]) -> None:
    tbl_pr = table._tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(USABLE_WIDTH_DXA))
    tbl_pr.append(width)

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:color"), "auto")
        element.set(qn("w:sz"), "4")
        borders.append(element)
    tbl_pr.append(borders)

    widths = column_widths(headers)
    grid = table._tbl.tblGrid
    if grid is not None:
        table._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width_dxa in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width_dxa))
        grid.append(col)
    table._tbl.insert(1, grid)

    for row_index, row in enumerate(table.rows):
        set_row_cant_split(row)
        if row_index == 0:
            set_row_repeats(row)
        for col_index, cell in enumerate(row.cells):
            set_cell_margins(cell)
            if col_index < len(widths):
                set_cell_width(cell, widths[col_index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(TABLE_FONT_PT)


def discover_styles(doc: Document):
    defined_ids = {style.style_id for style in doc.styles if style.style_id}
    level_to_id: dict[int, str] = {}
    list_id = "ListParagraph"
    for paragraph in doc.paragraphs:
        name = paragraph.style.name if paragraph.style else ""
        p_pr = paragraph._p.find(qn("w:pPr"))
        p_style = p_pr.find(qn("w:pStyle")) if p_pr is not None else None
        if p_style is None:
            continue
        style_id = p_style.get(qn("w:val"))
        if name.startswith("Heading"):
            level = name.replace("Heading", "").strip()
            if level.isdigit():
                level_to_id.setdefault(int(level), style_id)
        elif name == "List Paragraph":
            list_id = style_id
    title_id = "Title" if "Title" in defined_ids else level_to_id.get(1, "Heading1")
    return defined_ids, level_to_id, list_id, title_id


def style_for_heading(report_level: int, defined_ids: set[str], level_to_id: dict[int, str], title_id: str) -> str:
    if report_level == 1:
        return title_id
    template_level = report_level - 1
    fallback = f"Heading{template_level}"
    return level_to_id.get(template_level) or (fallback if fallback in defined_ids else level_to_id.get(1, "Heading1"))


def infer_title(markdown: str, fallback: str) -> str:
    for line in markdown.splitlines():
        match = re.match(r"#\s+(.+)", line.strip())
        if match:
            return strip_links(match.group(1)).strip()
    return fallback


def render_report(source_md: Path, output_docx: Path, header_title: str | None = None) -> None:
    markdown = source_md.read_text(encoding="utf-8")
    title = header_title or infer_title(markdown, output_docx.stem.replace("_", " ").title())

    doc = Document(str(TEMPLATE))
    defined_ids, level_to_id, list_id, title_id = discover_styles(doc)

    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child.tag in (qn("w:p"), qn("w:tbl")):
            body.remove(child)

    def emit_para():
        paragraph = doc.add_paragraph()
        sect_pr.addprevious(paragraph._p)
        return paragraph

    def emit_table(rows: int, cols: int):
        table = doc.add_table(rows=rows, cols=cols)
        sect_pr.addprevious(table._tbl)
        return table

    for block in parse_blocks(markdown):
        kind = block[0]
        if kind == "h":
            paragraph = emit_para()
            set_paragraph_style(paragraph, style_for_heading(block[1], defined_ids, level_to_id, title_id))
            add_runs(paragraph, inline_runs(block[2]))
        elif kind == "p":
            add_runs(emit_para(), inline_runs(block[1]))
        elif kind == "bullet":
            paragraph = emit_para()
            set_paragraph_style(paragraph, list_id)
            set_bullet(paragraph)
            add_runs(paragraph, inline_runs(block[1]))
        elif kind == "number":
            paragraph = emit_para()
            set_paragraph_style(paragraph, list_id)
            add_runs(paragraph, [(block[1] + ".  ", {})] + inline_runs(block[2]))
        elif kind == "quote":
            paragraph = emit_para()
            add_runs(paragraph, [(text, {**fmt, "italic": True}) for text, fmt in inline_runs(block[1])])
            paragraph.paragraph_format.left_indent = Pt(18)
        elif kind == "table":
            header, aligns, rows = block[1], block[2], block[3]
            cols = len(header)
            table = emit_table(len(rows) + 1, cols)
            for col, heading in enumerate(header):
                add_runs(table.cell(0, col).paragraphs[0], [(strip_links(heading), {"bold": True})])
            for row_index, row in enumerate(rows):
                for col in range(cols):
                    paragraph = table.cell(row_index + 1, col).paragraphs[0]
                    add_runs(paragraph, inline_runs(row[col] if col < len(row) else ""))
                    if aligns[col] == "right":
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    elif aligns[col] == "center":
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_table(table, header)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            for run in paragraph.runs:
                if HEADER_OLD in run.text:
                    run.text = run.text.replace(HEADER_OLD, title)

    doc.save(str(output_docx))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source_md", nargs="?", default="second_seed_report.md", help="Markdown report to render")
    parser.add_argument("output_docx", nargs="?", default="second_seed_report.docx", help="DOCX path to write")
    parser.add_argument("--title", help="Running-header title. Defaults to the Markdown H1.")
    args = parser.parse_args()

    source_md = (ROOT / args.source_md).resolve()
    output_docx = (ROOT / args.output_docx).resolve()
    render_report(source_md, output_docx, args.title)
    print(f"saved -> {output_docx.name}")


if __name__ == "__main__":
    main()
